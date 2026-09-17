from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from app.schemas.document import Block, BlockType, ParsedDocument
from app.services.styles import (
    add_bibliography_item,
    add_gost_paragraph,
    add_heading_center,
    setup_document_styles,
)

logger = logging.getLogger(__name__)


def render_document(
    parsed: ParsedDocument,
    output_path: str | Path,
    *,
    title_page_marker_skip: bool = True,
) -> str:
    """
    Собирает DOCX из ParsedDocument.

    :param parsed: разобранный документ
    :param output_path: куда сохранить
    :param title_page_marker_skip: если True — блоки из секции 'preamble'
                                   не рендерятся (титульник генерируется отдельно)
    """
    doc = Document()
    setup_document_styles(doc)

    for block in parsed.blocks:
        if title_page_marker_skip and block.section_id == "preamble":
            continue

        _render_block(doc, block, parsed)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info("Rendered document to %s (%d blocks)", output_path, len(parsed.blocks))
    return str(output_path)


def _render_block(doc: Document, block: Block, parsed: ParsedDocument) -> None:
    if block.type == BlockType.HEADING:
        add_heading_center(doc, block.text or "")

    elif block.type == BlockType.PARAGRAPH:
        add_gost_paragraph(doc, block.text or "")

    elif block.type == BlockType.TABLE:
        _render_table(doc, block)

    elif block.type == BlockType.IMAGE:
        _render_image(doc, block, parsed)

    else:
        logger.warning("Unknown block type: %s", block.type)


def _render_table(doc: Document, block: Block) -> None:
    rows = block.rows or []
    if not rows:
        return

    # Подпись сверху (ГОСТ 7.32 — над таблицей)
    if block.caption:
        caption = doc.add_paragraph(block.caption)
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT

    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"

    for i, row in enumerate(rows):
        for j in range(cols):
            value = row[j] if j < len(row) else ""
            table.cell(i, j).text = value


def _render_image(doc: Document, block: Block, parsed: ParsedDocument) -> None:
    image_path = block.image_path
    if not image_path:
        image_path = parsed.media.get(block.image_id or "", "")
    if not image_path or not Path(image_path).exists():
        logger.warning("Image not found for block %s: %s", block.id, image_path)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Cm(12))

    if block.caption:
        cap = doc.add_paragraph(block.caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER