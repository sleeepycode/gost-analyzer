from docx import Document
from app.schemas.blocks import ParagraphBlock, TableBlock, ImageBlock

def extract_metrics(source_path: str) -> dict:
    doc = Document(source_path)

    text_parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    full_text = '\n'.join(text_parts)
    text_length = len(full_text)
    tables_count = len(doc.tables)

    images_count = 0
    for rel in doc.part._rels.values():
        target_ref = getattr(rel, 'target_ref', '') or ''
        if 'image' in target_ref:
            images_count += 1

    return {
        'text_length': text_length,
        'tables_count': tables_count,
        'images_count': images_count,
        'paragraph_count': len(text_parts),
    }


def validate_source_document(source_path: str) -> dict:
    metrics = extract_metrics(source_path)
    errors: list[str] = []
    warnings: list[str] = []

    if metrics['text_length'] < 3000:
        errors.append('Документ невалиден: менее 3000 символов текста.')

    if metrics['tables_count'] + metrics['images_count'] < 1:
        errors.append('Документ невалиден: отсутствуют таблицы и рисунки.')

    if metrics['paragraph_count'] < 3:
        warnings.append('Документ содержит очень мало абзацев, проверьте корректность входного файла.')

    return {
        'is_valid': len(errors) == 0,
        'metrics': metrics,
        'errors': errors,
        'warnings': warnings,
    }

def validate_blocks(blocks: list) -> dict:
    text_parts = []
    tables_count = 0
    images_count = 0

    for block in blocks:
        if isinstance(block, ParagraphBlock):
            if block.text.strip():
                text_parts.append(block.text.strip())
        elif isinstance(block, TableBlock):
            tables_count += 1
        elif isinstance(block, ImageBlock):
            images_count += 1

    text_length = len("\n".join(text_parts))

    errors = []
    warnings = []

    if text_length < 3000:
        errors.append(f"Недостаточный объём текста: {text_length} символов, требуется минимум 3000.")

    if tables_count + images_count < 1:
        errors.append("Документ должен содержать хотя бы одну таблицу или один рисунок.")

    return {
        "is_valid": len(errors) == 0,
        "errors": errors,
        "warnings": warnings,
        "metrics": {
            "text_length": text_length,
            "tables_count": tables_count,
            "images_count": images_count,
        },
    }