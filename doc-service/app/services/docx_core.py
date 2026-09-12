import json
import zipfile
from pathlib import Path
from typing import List, Dict, Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from app.services.gost_applier import apply_gost_formatting

PROJECT_ROOT = Path('storage') / 'projects'
PROJECT_ROOT.mkdir(parents=True, exist_ok=True)


def ensure_project_dir(project_id: str) -> Path:
    project_dir = PROJECT_ROOT / project_id
    project_dir.mkdir(parents=True, exist_ok=True)
    return project_dir


def save_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def paragraph_contains_drawing(paragraph: Paragraph) -> bool:
    xml = paragraph._p.xml
    return '<w:drawing' in xml or '<pic:pic' in xml


def iter_block_items(parent: DocxDocument):
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _make_id(prefix: str, idx: int) -> str:
    return f'{prefix}_{idx + 1}'


def extract_images_from_docx(docx_path: str, media_dir: str) -> list[str]:
    media_paths: list[str] = []
    media_root = Path(media_dir)
    media_root.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(docx_path, 'r') as archive:
        for name in archive.namelist():
            if name.startswith('word/media/'):
                filename = Path(name).name
                target = media_root / filename
                with archive.open(name) as src, target.open('wb') as dst:
                    dst.write(src.read())
                media_paths.append(str(target))

    return media_paths


def extract_docx(docx_path: str, out_media_dir: str, project_id: str | None = None) -> Dict[str, Any]:
    doc = Document(docx_path)
    image_files = extract_images_from_docx(docx_path, out_media_dir)

    content_blocks = []
    paragraphs: List[Dict[str, Any]] = []
    tables: List[Dict[str, Any]] = []
    images: List[Dict[str, Any]] = []

    paragraph_count = 0
    table_count = 0
    image_index = 0

    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            has_image = paragraph_contains_drawing(item)

            if has_image and image_index < len(image_files):
                # Изображение внутри параграфа
                image_meta = {
                    'id': _make_id('image', image_index),
                    'path': image_files[image_index],
                    'source': 'original_docx_image',
                    'type': 'image',
                    'caption': None,
                    'position': image_index,
                    'context_text': text if text else None,
                    'insert_before_paragraph': paragraph_count  # <-- привязка
                }
                images.append(image_meta)
                content_blocks.append({
                    'type': 'image',
                    'data': image_meta
                })
                image_index += 1

            if text:
                para_meta = {
                    'id': _make_id('paragraph', paragraph_count),
                    'text': text,
                    'source': 'original_docx_paragraph',
                    'type': 'paragraph',
                    'position': paragraph_count,
                }
                paragraphs.append(para_meta)
                content_blocks.append({
                    'type': 'paragraph',
                    'data': para_meta
                })
                paragraph_count += 1

        elif isinstance(item, Table):
            table_data = []
            for row in item.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            table_meta = {
                'id': _make_id('table', table_count),
                'rows': table_data,
                'source': 'original_docx_table',
                'type': 'table',
                'position': table_count,
            }
            tables.append(table_meta)
            content_blocks.append({
                'type': 'table',
                'data': table_meta
            })
            table_count += 1

    # Добавляем оставшиеся изображения (если есть)
    while image_index < len(image_files):
        image_meta = {
            'id': _make_id('image', image_index),
            'path': image_files[image_index],
            'source': 'original_docx_image',
            'type': 'image',
            'caption': None,
            'position': image_index,
            'context_text': None,
            'insert_before_paragraph': None
        }
        images.append(image_meta)
        content_blocks.append({
            'type': 'image',
            'data': image_meta
        })
        image_index += 1

    result = {
        'paragraphs': paragraphs,
        'tables': tables,
        'images': images,
        'content_blocks': content_blocks
    }

    if project_id:
        project_dir = ensure_project_dir(project_id)
        save_json(project_dir / 'extracted.json', result)

    return result


def _is_section_heading(text: str) -> bool:
    if not text:
        return False
    normalized = text.strip()
    if len(normalized) > 120:
        return False
    lower = normalized.lower()
    headings = [
        'введение', 'теоретическая часть', 'теория', 'экспериментальная часть',
        'практическая часть', 'выводы', 'заключение', 'список литературы',
        'литература', 'результаты', 'анализ', 'обсуждение', 'методика',
        'цель', 'задачи'
    ]
    if any(lower.startswith(keyword) for keyword in headings):
        return True
    if ':' in normalized and len(normalized) < 90:
        return True
    if normalized.isupper() and len(normalized.split()) < 10:
        return True
    return False


SECTION_ORDER = [
    'title',
    'introduction',
    'practice',
    'conclusion',
    'bibliography',
]

SECTION_TITLES = {
    'title': None,
    'introduction': 'Введение',
    'practice': 'Практическая часть',
    'conclusion': 'Заключение',
    'bibliography': 'Список литературы',
}

SECTION_ALIASES = {
    'theory': 'introduction',
    'introduction': 'introduction',
    'practice': 'practice',
    'conclusion': 'conclusion',
    'references': 'bibliography',
    'bibliography': 'bibliography',
    'title': 'title',
}


def _normalize_section_key(section_name: str) -> str | None:
    if not section_name:
        return None
    key = section_name.strip().lower()
    return SECTION_ALIASES.get(key)


def _get_paragraph_text(item: Any) -> str:
    if isinstance(item, dict):
        return str(item.get('text', ''))
    return str(item)


def _add_source_note(doc: Document, source_text: str) -> None:
    note = doc.add_paragraph()
    run = note.add_run(f'Источник: {source_text}')
    run.italic = True
    note.paragraph_format.line_spacing = 1.0
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(0)


def _insert_image(doc: Document, img_path: str, caption: str | None = None, source: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    try:
        paragraph.add_run().add_picture(img_path, width=Cm(15))
    except Exception:
        paragraph.add_run('[Не удалось вставить изображение]')
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        cap.style = 'Caption'
    if source:
        _add_source_note(doc, source)


def _add_table(doc: Document, table_meta: Dict[str, Any]) -> None:
    rows = table_meta.get('rows', [])
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]) if rows else 0)
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text
    if table_meta.get('source'):
        _add_source_note(doc, table_meta['source'])


def _add_paragraph_or_heading(doc: Document, text: str) -> None:
    if _is_section_heading(text):
        doc.add_heading(text, level=1)
    else:
        doc.add_paragraph(text)


def _append_generated_sections(doc: Document, ml_response: Dict[str, Any]) -> None:
    for section in ml_response.get('generated_sections', []):
        title = section.get('title') or section.get('section') or 'Раздел'
        text = section.get('text', '')
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)

    for suggestion in ml_response.get('content_suggestions', []):
        if suggestion.get('action') == 'generate':
            title = suggestion.get('title') or suggestion.get('target', 'Раздел')
            message = suggestion.get('message', '')
            doc.add_heading(title, level=1)
            doc.add_paragraph(message)

    bibliography = ml_response.get('bibliography', [])
    if bibliography:
        doc.add_heading('Список литературы', level=1)
        for ref in bibliography:
            p = doc.add_paragraph(ref)
            p.style = 'List Number'

    gost_validation = ml_response.get('gost_validation')
    if gost_validation and gost_validation.get('issues'):
        doc.add_heading('Рекомендации по ГОСТ', level=2)
        for issue in gost_validation.get('issues', []):
            doc.add_paragraph(f'- {issue}')


def assemble_structure(
    structure: Dict[str, Any],
    output_path: str,
    project_id: str | None = None,
    ml_response: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    tmp_path = str(Path(output_path).with_suffix('.tmp.docx'))
    doc = Document()

    section_keys = [key for key in SECTION_ORDER if key in structure]
    if section_keys:
        for key in SECTION_ORDER:
            if key not in structure:
                continue

            if key == 'bibliography':
                bibliography = structure.get('bibliography') or []
                if bibliography:
                    doc.add_heading(SECTION_TITLES[key], level=1)
                    for ref in bibliography:
                        p = doc.add_paragraph(str(ref))
                        p.style = 'List Number'
                continue

            section_data = structure.get(key)
            if section_data is None:
                continue

            if SECTION_TITLES[key]:
                doc.add_heading(SECTION_TITLES[key], level=1)

            if isinstance(section_data, str):
                _add_paragraph_or_heading(doc, section_data)
            elif isinstance(section_data, list):
                for item in section_data:
                    _add_paragraph_or_heading(doc, _get_paragraph_text(item))
            elif isinstance(section_data, dict):
                _add_paragraph_or_heading(doc, _get_paragraph_text(section_data))

        # Добавляем дополнительные содержимое, если есть
        paragraphs = structure.get('paragraphs', []) or []
        images = structure.get('images', []) or []
        tables = structure.get('tables', []) or []

        for idx, item in enumerate(paragraphs):
            _add_paragraph_or_heading(doc, _get_paragraph_text(item))
            for img in images:
                placement = img.get('placement', {})
                if placement.get('paragraph_index') == idx:
                    _insert_image(doc, img.get('path', ''), img.get('caption'), img.get('source'))

        for table_meta in tables:
            _add_table(doc, table_meta)
    else:
        paragraphs = structure.get('paragraphs', []) or []
        images = structure.get('images', []) or []
        tables = structure.get('tables', []) or []

        for idx, item in enumerate(paragraphs):
            _add_paragraph_or_heading(doc, _get_paragraph_text(item))
            for img in images:
                placement = img.get('placement', {})
                if placement.get('paragraph_index') == idx:
                    _insert_image(doc, img.get('path', ''), img.get('caption'), img.get('source'))

        for table_meta in tables:
            _add_table(doc, table_meta)

    if ml_response:
        _append_generated_sections(doc, ml_response)
        if project_id:
            project_dir = ensure_project_dir(project_id)
            save_json(project_dir / 'ml_response.json', ml_response)

    doc.save(tmp_path)

    try:
        apply_gost_formatting(tmp_path, output_path)
        Path(tmp_path).unlink(missing_ok=True)
        return {'status': 'completed', 'output_path': output_path}
    except Exception as e:
        return {'status': 'failed', 'error': str(e)}


def analyze_project_stub(document_text: str, image_paths: List[str], topic: str) -> Dict[str, Any]:
    return {
        'generated_sections': [
            {'title': f'Автосгенерированное: {topic}', 'text': 'Это пример сгенерированного раздела.'}
        ],
        'bibliography': [],
        'images': [],
        'meta': {'note': 'ml stub'}
    }
