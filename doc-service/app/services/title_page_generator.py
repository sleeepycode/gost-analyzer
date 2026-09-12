from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from app.core.config import settings
from app.services.extractor import iter_block_items


TITLE_KEYWORDS = [
    "министерство", "университет", "институт", "кафедра", "факультет",
    "лабораторная работа", "выполнил", "проверил", "руководитель", "москва"
]

BODY_START_KEYWORDS = [
    "введение", "цель работы", "цель лабораторной работы", "ход работы",
    "теоретические сведения", "практическая часть", "выполнение работы", "заключение"
]


def normalize_text(text: str) -> str:
    return " ".join(text.lower().strip().split())


def is_title_like(text: str) -> bool:
    t = normalize_text(text)
    return any(keyword in t for keyword in TITLE_KEYWORDS)


def is_body_start(text: str) -> bool:
    t = normalize_text(text)
    return any(t.startswith(keyword) for keyword in BODY_START_KEYWORDS)


def _find_body_start_index(blocks: list) -> int | None:
    for i, block in enumerate(blocks[:40]):
        if hasattr(block, 'text') and block.text.strip():
            if is_body_start(block.text.strip()):
                return i
    return None


def remove_existing_title_page(doc: Document) -> None:
    """Удаляет существующий титульный лист из документа"""
    blocks = list(iter_block_items(doc))
    body_start_index = _find_body_start_index(blocks)
    
    if body_start_index is not None and body_start_index > 0:
        body = doc.element.body
        children = list(body)
        for i in range(body_start_index):
            body.remove(children[i])


def generate_title_page(
    input_path: str,
    output_path: str,
    department: str,
    discipline: str,
    lab_number: str,
    lab_title: str,
    student_group: str,
    student_name: str,
    reviewer_name: str,
) -> None:
    """
    Генерирует титульный лист по каноничному формату
    """
    doc = Document(input_path)

    remove_existing_title_page(doc)

    first_paragraph = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()

    def insert_title_paragraph(text: str, alignment: WD_ALIGN_PARAGRAPH, size: int | float = None,
                                  bold: bool = False, space_before: float = 0, space_after: float = 0):
        paragraph = first_paragraph.insert_paragraph_before(text)
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(space_before)
        paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.line_spacing = 1.5
        for run in paragraph.runs:
            run.font.name = settings.font_name
            run.font.size = Pt(size or settings.font_size_pt)
            run.bold = bold
        return paragraph

    title_lines = [
        (f"{settings.year}", WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        (f"{settings.city},", WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, None, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, None, False, 0, 0),
        (f"{reviewer_name}", WD_ALIGN_PARAGRAPH.RIGHT, settings.font_size_pt, False, 0, 0),
        (f"Проверил:", WD_ALIGN_PARAGRAPH.RIGHT, settings.font_size_pt, False, 0, 0),
        (f"{student_name}", WD_ALIGN_PARAGRAPH.RIGHT, settings.font_size_pt, False, 0, 0),
        (f"Выполнил: студент группы {student_group}", WD_ALIGN_PARAGRAPH.RIGHT, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        (f'«{lab_title}»', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('на тему:', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        (f"по дисциплине «{discipline}»", WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        (f"Отчет по лабораторной работе №{lab_number}", WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        (f"Кафедра «{department}»", WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, False, 0, 0),
        ('«Московский технический университет связи и информатики»', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('образовательное учреждение высшего образования', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('бюджетное', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('Ордена трудового Красного Знамени федеральное государственное', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('МАССОВЫХ КОММУНИКАЦИЙ РОССИЙСКОЙ ФЕДЕРАЦИИ', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0),
        ('МИНИСТЕРСТВО ЦИФРОВОГО РАЗВИТИЯ, СВЯЗИ И', WD_ALIGN_PARAGRAPH.CENTER, settings.font_size_pt, True, 0, 0)
    ]

    for text, alignment, size, bold, space_before, space_after in reversed(title_lines):
        insert_title_paragraph(text, alignment, size, bold, space_before, space_after)

    page_break_paragraph = first_paragraph.insert_paragraph_before('')
    page_break_run = page_break_paragraph.add_run()
    page_break_run.add_break(WD_BREAK.PAGE)

    doc.save(output_path)