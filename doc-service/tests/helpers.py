"""Генераторы тестовых DOCX. Используют python-docx напрямую."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches


def make_simple_docx(path: Path) -> Path:
    """Простой документ: три параграфа, без таблиц и картинок."""
    doc = Document()
    doc.add_paragraph("Это первый параграф тестового документа.")
    doc.add_paragraph("Это второй параграф с обычным текстом.")
    doc.add_paragraph("Это третий параграф для проверки парсера.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def make_with_table_docx(path: Path) -> Path:
    """Документ с одной таблицей 3x3."""
    doc = Document()
    doc.add_paragraph("Введение")
    doc.add_paragraph("Текст до таблицы.")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    data = [
        ["Заголовок 1", "Заголовок 2", "Заголовок 3"],
        ["a", "b", "c"],
        ["d", "e", "f"],
    ]
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            table.cell(i, j).text = value

    doc.add_paragraph("Текст после таблицы.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def make_with_images_docx(path: Path) -> Path:
    """Документ с двумя картинками."""
    from PIL import Image

    tmp_dir = path.parent
    tmp_dir.mkdir(parents=True, exist_ok=True)

    img1_path = tmp_dir / "_tmp_img1.png"
    img2_path = tmp_dir / "_tmp_img2.png"

    Image.new("RGB", (50, 50), color="red").save(img1_path)
    Image.new("RGB", (50, 50), color="blue").save(img2_path)

    doc = Document()
    doc.add_paragraph("Первый параграф перед картинкой.")

    p1 = doc.add_paragraph()
    p1.add_run().add_picture(str(img1_path), width=Inches(1))

    doc.add_paragraph("Параграф между картинками.")

    p2 = doc.add_paragraph()
    p2.add_run().add_picture(str(img2_path), width=Inches(1))

    doc.add_paragraph("Параграф после картинок.")

    doc.save(str(path))

    img1_path.unlink(missing_ok=True)
    img2_path.unlink(missing_ok=True)

    return path


def make_with_sections_docx(path: Path) -> Path:
    """Документ с каноничными разделами."""
    doc = Document()
    doc.add_paragraph("Введение")
    doc.add_paragraph("Цель работы — проверить парсер.")

    doc.add_paragraph("Теоретическая часть")
    doc.add_paragraph("Немного теории.")

    doc.add_paragraph("Практическая часть")
    doc.add_paragraph("Ход работы и результаты.")

    doc.add_paragraph("Заключение")
    doc.add_paragraph("Выводы по работе.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def make_mixed_docx(path: Path) -> Path:
    """Всё вместе: заголовки, текст, таблица, картинка."""
    from PIL import Image

    tmp_dir = path.parent
    tmp_dir.mkdir(parents=True, exist_ok=True)
    img_path = tmp_dir / "_tmp_mixed.png"
    Image.new("RGB", (60, 60), color="green").save(img_path)

    doc = Document()
    doc.add_paragraph("Введение")
    doc.add_paragraph("Короткое введение.")

    doc.add_paragraph("Практическая часть")
    doc.add_paragraph("Первый абзац практики.")

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Параметр"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "x"
    table.cell(1, 1).text = "42"

    doc.add_paragraph("Абзац после таблицы.")

    p = doc.add_paragraph()
    p.add_run().add_picture(str(img_path), width=Inches(1))

    doc.add_paragraph("Заключение")
    doc.add_paragraph("Финальные выводы.")

    doc.save(str(path))
    img_path.unlink(missing_ok=True)
    return path