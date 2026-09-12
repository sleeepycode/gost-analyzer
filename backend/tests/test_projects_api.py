import json
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("test")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _valid_processing_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Тестовый текст лабораторной работы.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _client(tmp_path):
    settings.storage_dir = str(tmp_path / "storage")
    settings.input_dir = str(tmp_path / "storage" / "inputs")
    settings.output_dir = str(tmp_path / "storage" / "outputs")
    settings.report_dir = str(tmp_path / "storage" / "reports")
    settings.projects_dir = str(tmp_path / "storage" / "projects")

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app)


def test_upload_project_file_creates_project_structure(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "uploaded"
        project_id = body["project_id"]

        project_dir = Path(settings.projects_dir) / project_id
        assert (project_dir / "input").exists()
        assert (project_dir / "images").exists()
        assert (project_dir / "output").exists()
        metadata_path = project_dir / "metadata.json"
        assert metadata_path.exists()

        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        assert metadata["project_id"] == project_id
        assert metadata["status"] == "uploaded"
        assert metadata["source_filename"] == "source.docx"
        assert metadata["input_file"]
        assert metadata["images"] == []
        assert metadata["output_file"] is None
        assert metadata["ml_result"] is None
        assert metadata["errors"] == []


def test_upload_project_file_rejects_unsupported_extension(tmp_path):
    with _client(tmp_path) as client:
        files = {"file": ("notes.txt", b"text", "text/plain")}
        resp = client.post("/projects/upload", files=files)
        assert resp.status_code == 400
        body = resp.json()
        assert body["code"] == "invalid_file_format"
        assert "DOCX" in body["message"]


def test_download_project_result_returns_latest_completed_output(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        task_data = {
            "user_id": "user-1",
            "project_id": project_id,
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        task_files = {
            "file": (
                "valid.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        task_resp = client.post("/tasks", data=task_data, files=task_files)
        assert task_resp.status_code == 200

        download_pdf = client.get(f"/projects/{project_id}/download?user_id=user-1&format=pdf")
        assert download_pdf.status_code == 200
        assert download_pdf.headers["content-type"] == "application/pdf"

        download_docx = client.get(f"/projects/{project_id}/download?user_id=user-1&format=docx")
        assert download_docx.status_code == 200
        assert (
            download_docx.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


def test_download_project_result_forbidden_for_other_user(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        forbidden_resp = client.get(f"/projects/{project_id}/download?user_id=user-2")
        assert forbidden_resp.status_code == 403


def test_delete_project_removes_project_tasks_and_project_folder(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        task_data = {
            "user_id": "user-1",
            "project_id": project_id,
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        task_files = {
            "file": (
                "valid.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        task_resp = client.post("/tasks", data=task_data, files=task_files)
        assert task_resp.status_code == 200
        task_id = task_resp.json()["task_id"]

        delete_resp = client.delete(f"/projects/{project_id}?user_id=user-1")
        assert delete_resp.status_code == 200
        assert delete_resp.json()["status"] == "deleted"

        status_resp = client.get(f"/projects/{project_id}")
        assert status_resp.status_code == 404

        task_status_resp = client.get(f"/tasks/{task_id}")
        assert task_status_resp.status_code == 404

        project_dir = Path(settings.projects_dir) / project_id
        assert not project_dir.exists()


def test_delete_project_forbidden_for_other_user(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        forbidden_resp = client.delete(f"/projects/{project_id}?user_id=user-2")
        assert forbidden_resp.status_code == 403


def test_project_analyze_and_get_analysis(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        analyze_resp = client.post(f"/projects/{project_id}/analyze?user_id=user-1")
        assert analyze_resp.status_code == 200
        analyze_body = analyze_resp.json()
        assert analyze_body["status"] == "ready"
        assert analyze_body["analysis"]["project_id"] == project_id
        assert "preview" in analyze_body["analysis"]

        get_resp = client.get(f"/projects/{project_id}/analysis?user_id=user-1")
        assert get_resp.status_code == 200
        assert get_resp.json()["analysis"]["project_id"] == project_id

        metadata_path = Path(settings.projects_dir) / project_id / "metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        assert metadata["analysis"]["project_id"] == project_id


def test_upload_additional_files_to_existing_project_separates_input_and_images(tmp_path):
    with _client(tmp_path) as client:
        create_files = {
            "file": (
                "base.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        create_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=create_files)
        assert create_resp.status_code == 200
        project_id = create_resp.json()["project_id"]

        image_files = {"file": ("scan.jpg", b"jpeg-bytes", "image/jpeg")}
        add_image_resp = client.post(
            f"/projects/{project_id}/files",
            data={"user_id": "user-1"},
            files=image_files,
        )
        assert add_image_resp.status_code == 200
        assert add_image_resp.json()["file_type"] == "image"

        project_dir = Path(settings.projects_dir) / project_id
        assert (project_dir / "images" / "scan.jpg").exists()

        metadata_path = project_dir / "metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        assert metadata["source_filename"] == "base.docx"
        assert "scan.jpg" in metadata["images"]

        suggestions = metadata.get("image_suggestions") or []
        assert len(suggestions) == 1
        assert suggestions[0]["ocr_text"] == "U=IR"
        assert suggestions[0]["image_name"] == "scan.jpg"
        assert "scan.jpg" in metadata.get("image_ml_results", {})


def test_post_process_project_runs_same_pipeline_as_tasks(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        form = {
            "user_id": "user-1",
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        proc_resp = client.post(f"/projects/{project_id}/process", data=form)
        assert proc_resp.status_code == 200
        body = proc_resp.json()
        assert body["status"] == "ready"
        task_id = body["task_id"]

        out_pdf = Path(settings.projects_dir) / project_id / "output" / f"{task_id}.pdf"
        out_docx = Path(settings.projects_dir) / project_id / "output" / f"{task_id}.docx"
        assert out_pdf.exists()
        assert out_docx.exists()


def test_suggestions_get_and_apply(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        client.post(
            f"/projects/{project_id}/files",
            data={"user_id": "user-1"},
            files={"file": ("fig1.png", b"\x89PNG\r\n", "image/png")},
        )

        form = {
            "user_id": "user-1",
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        assert client.post(f"/projects/{project_id}/process", data=form).status_code == 200

        sug_resp = client.get(f"/projects/{project_id}/suggestions?user_id=user-1")
        assert sug_resp.status_code == 200
        items = sug_resp.json()["suggestions"]
        assert len(items) >= 1
        sid = items[0]["id"]

        apply_resp = client.post(
            f"/projects/{project_id}/suggestions/apply?user_id=user-1",
            json={"suggestion_ids": [sid]},
        )
        assert apply_resp.status_code == 200
        assert sid in apply_resp.json()["applied_ids"]

        meta = json.loads(
            (Path(settings.projects_dir) / project_id / "metadata.json").read_text(encoding="utf-8")
        )
        applied = [s for s in meta["image_suggestions"] if s.get("id") == sid]
        assert applied and applied[0].get("applied") is True
