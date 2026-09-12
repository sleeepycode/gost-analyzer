from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app


def _error_message(resp) -> str:
    body = resp.json()
    if "message" in body:
        return str(body["message"])
    return str(body.get("detail", ""))


@pytest.fixture()
def client(tmp_path):
    settings.storage_dir = str(tmp_path / "storage")
    settings.input_dir = str(tmp_path / "storage" / "inputs")
    settings.output_dir = str(tmp_path / "storage" / "outputs")
    settings.report_dir = str(tmp_path / "storage" / "reports")
    settings.projects_dir = str(tmp_path / "storage" / "projects")

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _form_data():
    return {
        "user_id": "demo-user-1",
        "faculty": "ФКТ",
        "department": "Кафедра ИС",
        "student_group": "БПИ-01",
        "lab_title": "Тестирование API",
        "lab_number": "1",
        "student_name": "Иванов И.И.",
        "reviewer_name": "Петров П.П.",
        "discipline": "Программирование",
    }


def _valid_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Тестовый текст лабораторной работы.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _invalid_by_rules_docx_bytes() -> bytes:
    doc = Document()
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_create_task_success_and_artifacts_available(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "completed"
    task_id = data["task_id"]

    status_resp = client.get(f"/tasks/{task_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["has_output"] is True
    assert status_data["has_report"] is True

    download_pdf = client.get(f"/tasks/{task_id}/download?user_id=demo-user-1&format=pdf")
    assert download_pdf.status_code == 200
    assert download_pdf.headers["content-type"] == "application/pdf"

    download_docx = client.get(f"/tasks/{task_id}/download?user_id=demo-user-1&format=docx")
    assert download_docx.status_code == 200
    assert (
        download_docx.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    report_resp = client.get(f"/tasks/{task_id}/report?user_id=demo-user-1")
    assert report_resp.status_code == 200
    assert report_resp.headers["content-type"].startswith("application/json")


def test_create_task_rejects_non_docx(client: TestClient):
    files = {
        "file": (
            "bad.txt",
            b"plain text",
            "text/plain",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    body = resp.json()
    assert body.get("code") == "invalid_file_format"
    assert "DOCX" in body.get("message", "")


def test_create_task_rejects_empty_docx(client: TestClient):
    files = {
        "file": (
            "empty.docx",
            b"",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    assert "пустой" in _error_message(resp).lower()


def test_create_task_rejects_corrupted_docx(client: TestClient):
    files = {
        "file": (
            "corrupted.docx",
            b"this is not a real docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    assert "docx" in _error_message(resp).lower()


def test_create_task_fails_business_validation(client: TestClient):
    files = {
        "file": (
            "too_short.docx",
            _invalid_by_rules_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "failed"
    assert data["report"]["status"] == "failed"
    assert data["report"]["errors"]


def test_get_unknown_task_returns_404(client: TestClient):
    resp = client.get("/tasks/not-existing-id")
    assert resp.status_code == 404


def test_list_tasks_history_for_frontend(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    create_resp = client.post("/tasks", data=_form_data(), files=files)
    assert create_resp.status_code == 200
    created_task_id = create_resp.json()["task_id"]

    history_resp = client.get("/tasks?limit=10")
    assert history_resp.status_code == 200
    body = history_resp.json()
    assert "items" in body
    assert isinstance(body["items"], list)
    assert len(body["items"]) >= 1

    first = body["items"][0]
    assert first["task_id"] == created_task_id
    assert "status" in first
    assert "user_id" in first
    assert "original_filename" in first
    assert "created_at" in first
    assert "has_output" in first
    assert "has_report" in first


def test_list_tasks_can_filter_by_user_id(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    user1_data = _form_data()
    user1_data["user_id"] = "user-1"
    user2_data = _form_data()
    user2_data["user_id"] = "user-2"

    resp1 = client.post("/tasks", data=user1_data, files=files)
    assert resp1.status_code == 200
    task1 = resp1.json()["task_id"]

    resp2 = client.post("/tasks", data=user2_data, files=files)
    assert resp2.status_code == 200
    task2 = resp2.json()["task_id"]

    filtered_user1 = client.get("/tasks?limit=10&user_id=user-1")
    assert filtered_user1.status_code == 200
    items1 = filtered_user1.json()["items"]
    assert items1
    assert all(item["user_id"] == "user-1" for item in items1)
    assert any(item["task_id"] == task1 for item in items1)
    assert not any(item["task_id"] == task2 for item in items1)


def test_delete_task_removes_db_record_and_files(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    create_resp = client.post("/tasks", data=_form_data(), files=files)
    assert create_resp.status_code == 200
    task_id = create_resp.json()["task_id"]

    status_before = client.get(f"/tasks/{task_id}")
    assert status_before.status_code == 200
    assert status_before.json()["has_output"] is True
    assert status_before.json()["has_report"] is True

    delete_resp = client.delete(f"/tasks/{task_id}?user_id=demo-user-1")
    assert delete_resp.status_code == 200
    delete_data = delete_resp.json()
    assert delete_data["task_id"] == task_id
    assert delete_data["status"] == "deleted"

    status_after = client.get(f"/tasks/{task_id}")
    assert status_after.status_code == 404


def test_delete_unknown_task_returns_404(client: TestClient):
    resp = client.delete("/tasks/not-existing-id?user_id=demo-user-1")
    assert resp.status_code == 404


def test_delete_task_forbidden_for_other_user(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    create_resp = client.post("/tasks", data=_form_data(), files=files)
    assert create_resp.status_code == 200
    task_id = create_resp.json()["task_id"]

    delete_resp = client.delete(f"/tasks/{task_id}?user_id=another-user")
    assert delete_resp.status_code == 403
    assert "другого пользователя" in _error_message(delete_resp).lower()


def test_download_and_report_forbidden_for_other_user(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    create_resp = client.post("/tasks", data=_form_data(), files=files)
    assert create_resp.status_code == 200
    task_id = create_resp.json()["task_id"]

    download_resp = client.get(f"/tasks/{task_id}/download?user_id=another-user")
    assert download_resp.status_code == 403

    report_resp = client.get(f"/tasks/{task_id}/report?user_id=another-user")
    assert report_resp.status_code == 403


def test_task_can_be_linked_to_project_and_update_project_status(client: TestClient):
    upload_files = {
        "file": (
            "source.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    upload_resp = client.post("/projects/upload", data={"user_id": "demo-user-1"}, files=upload_files)
    assert upload_resp.status_code == 200
    project_id = upload_resp.json()["project_id"]

    task_files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    task_data = _form_data()
    task_data["project_id"] = project_id
    task_resp = client.post("/tasks", data=task_data, files=task_files)
    assert task_resp.status_code == 200

    task_id = task_resp.json()["task_id"]
    task_status_resp = client.get(f"/tasks/{task_id}")
    assert task_status_resp.status_code == 200
    assert task_status_resp.json()["project_id"] == project_id

    project_status_resp = client.get(f"/projects/{project_id}")
    assert project_status_resp.status_code == 200
    assert project_status_resp.json()["status"] == "ready"

    assert (Path(settings.projects_dir) / project_id / "output" / f"{task_id}.pdf").exists()
    assert (Path(settings.projects_dir) / project_id / "output" / f"{task_id}.docx").exists()


def test_link_project_requires_user_id(client: TestClient):
    upload_files = {
        "file": (
            "source.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    upload_resp = client.post("/projects/upload", data={"user_id": "demo-user-1"}, files=upload_files)
    assert upload_resp.status_code == 200
    project_id = upload_resp.json()["project_id"]

    task_files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    task_data = _form_data()
    task_data.pop("user_id")
    task_data["project_id"] = project_id
    task_resp = client.post("/tasks", data=task_data, files=task_files)
    assert task_resp.status_code == 400
    assert "требуется user_id" in _error_message(task_resp).lower()
