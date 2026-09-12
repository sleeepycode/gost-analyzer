"""Проверка входного DOCX перед отправкой в doc-service (без правки файла)."""

from docx import Document


def _collect_text_parts(doc: Document) -> list[str]:
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    return parts


def extract_metrics(source_path: str) -> dict:
    doc = Document(source_path)
    text_parts = _collect_text_parts(doc)
    full_text = "\n".join(text_parts)

    images_count = 0
    for rel in doc.part._rels.values():
        target_ref = getattr(rel, "target_ref", "") or ""
        if "image" in target_ref:
            images_count += 1

    return {
        "text_length": len(full_text),
        "tables_count": len(doc.tables),
        "images_count": images_count,
        "paragraph_count": len(text_parts),
    }


def _is_content_valid(metrics: dict) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if metrics["text_length"] == 0:
        errors.append(
            "В файле обязателен текст. Допускаются также таблицы, изображения и другие элементы."
        )
    return len(errors) == 0, errors


def validate_source_document(source_path: str) -> dict:
    metrics = extract_metrics(source_path)
    errors: list[str] = []
    warnings: list[str] = []

    valid, content_errors = _is_content_valid(metrics)
    errors.extend(content_errors)

    if metrics["paragraph_count"] < 3 and metrics["text_length"] > 0:
        warnings.append("Документ содержит очень мало текста, проверьте корректность входного файла.")

    return {
        "is_valid": valid and len(errors) == 0,
        "metrics": metrics,
        "errors": errors,
        "warnings": warnings,
    }
